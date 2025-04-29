from kivymd.app import MDApp
from kivy.lang import Builder
from kivy.animation import Animation
from kivy.core.window import Window
from kivy.uix.screenmanager import Screen, ScreenManager
from kivy.clock import Clock
from docx import Document
import sqlite3

class Main(Screen):
    def generate(self):
        self.ids.cv.color = "#8D9A92FF"
        self.ids.cv.text = "generating..."
        
        Clock.schedule_once(self.generator, 1)

    def generator(self, *args):
        self.ids.cv.text = "generated :)"
        self.ids.cv.color = "#6AE394FF"
        
        # acessing input
        
        workplace = self.ids.work.text
        work_Achivement = self.ids.aone.text
        work_Achivement2 = self.ids.atwo.text
        work_Achivement3 = self.ids.athree.text
        professional_summary = self.ids.summary.text
        languages = self.ids.lang.text
        frameworks_and_libraries = self.ids.frame.text
        tools = self.ids.tool.text
        other = self.ids.other.text
        project_one = self.ids.proone.text
        project_two = self.ids.protwo.text
        project_three = self.ids.prothree.text
        ref = self.ids.ref.text
        portfolio_web = self.ids.pro.text

        #  Database stores past records of previous cvs generated 
        conn = sqlite3.connect("CV_History_Data.db")
        cursor = conn.cursor()

        cursor.execute('''
             CREATE TABLE IF NOT EXISTS details (
                 id INTEGER PRIMARY KEY,
                 work_place TEXT,
                 work_detail1 TEXT,
                 work_detail2 TEXT,
                 work_detail3 TEXT,
                 professional_summary TEXT,
                 languages TEXT,
                 frameworks_and_libraries TEXT,
                 tools TEXT,
                 Others TEXT,
                 ref TEXT,
                 project1 TEXT,
                 project2 TEXT,
                 project3 TEXT
                
                
                
             )
         ''')

    
        cursor.execute("""INSERT INTO details (
                 work_place,
                 work_detail1,
                 work_detail2,
                 work_detail3,
                 professional_summary ,
                 languages ,
                 frameworks_and_libraries ,
                tools ,
                 Others ,
                 ref,
                 project1,
                 project2,
                 project3) VALUES (?, ?, ?,?,?,?,?,?,?,?,?,?,?)""", 
                        (workplace,work_Achivement,work_Achivement2,work_Achivement3,professional_summary,languages,frameworks_and_libraries,tools,other,ref,project_one,project_two,project_three))
        

    
        conn.commit()
        # Generating Word document
        doc = Document()
        doc.add_heading(level=1, text="Andrew Muyunda")
        doc.add_paragraph("Kitwe, Zambia")
        doc.add_paragraph("+260777290515 | andrewmmuyunda777@gmail.com")
        doc.add_heading(level=0, text="")
        if professional_summary is "":
            summary = doc.add_paragraph("As a seasoned web and automation Developer with a proven track record of 2 years, "
            "I bring a wealth of expertise in crafting innovative and user-friendly web and desktop applications. "
            "Specializing in Python, Typescript, Javascript, HTML, CSS, MySQL, React, Next.js, and Kivy, I have "
            "successfully delivered various projects, showcasing my proficiency in web and desktop technologies. "
            "My commitment to staying abreast of industry trends and adopting best practices has allowed me to "
            "contribute significantly to project success. With a keen eye for detail and a passion for creating "
            "seamless user experiences, I am eager to leverage my skills and experience to drive excellence in "
            "future development endeavors.")
            
        else:
            doc.add_paragraph(professional_summary)    
        doc.add_heading(level=0, text="")

        skills = doc.add_heading("Skills:", 3)
        
        if languages is not "":
            doc.add_paragraph(style="List Bullet", text=languages)
        if frameworks_and_libraries is not "":
            doc.add_paragraph(style="List Bullet", text=frameworks_and_libraries)
        if tools is not "":
            doc.add_paragraph(style="List Bullet", text=tools)
        if other is not "":
            doc.add_paragraph(style="List Bullet", text="")                
        doc.add_heading(level=0, text="")
        work = doc.add_heading("Work experience:", 3)
        if workplace is not "":
            doc.add_paragraph( text=workplace)
            doc.add_paragraph(style="List Bullet", text=work_Achivement)
            doc.add_paragraph(style="List Bullet", text=work_Achivement2)
            doc.add_paragraph(style="List Bullet", text=work_Achivement2)
        project = doc.add_heading("personal projects:", 3)
        if project_one is "":
            doc.add_paragraph(style="List Bullet" , text=" Huddle - is a platform dedicated to helping you achieve your self-improvement goals. Here, you'll find all the tools and resources you need on your journey toward personal growth and self-development, all in one cozy place and also seeks to be a safe space for people to grow. ")
            doc.add_paragraph(style="List Bullet" , text="Autopal - is a platform that teaches you the fundamentals of automation and key concepts from web scraping to application development all in a detailed consise and structured manner")
            doc.add_paragraph(style="List Bullet" ,text="Tag - is my personalized Rag assistant")
        else:
            doc.add_paragraph(style="List Bullet",text=project_one)
            doc.add_paragraph(style="List Bullet",text=project_two)
            doc.add_paragraph(style="List Bullet",text=project_three)
        doc.add_heading("Education:", 3)
        doc.add_paragraph(style="List Bullet", text="Graduated from Mulungushi University")
        doc.add_paragraph("Bachelors in Science in Computer Science | 2022-2026")
        doc.add_paragraph(style="List Bullet", text="Completed secondary education at Ndeke Secondary School")
        doc.add_paragraph("G12 Certificate obtained | 2016-2021")
        doc.add_heading(level=0, text="")

        reference = doc.add_heading("Reference:", 3)
        if ref is "":
            doc.add_paragraph("Available on request")
        else :
            doc.add_paragraph(style="List Bullet" , text=ref)    
        if portfolio_web is not "":
            doc.add_heading(level=3,text="Portfolio Website")
            doc.add_paragraph(style="List Bullet",text=portfolio_web)
        
        if professional_summary or portfolio_web is not "":
            doc.save("C:\\Users\\AUTHENTIC PLUS STORE\\Desktop\\my-updatedf-cv.docx")
        else:
            doc.save("C:\\Users\\AUTHENTIC PLUS STORE\\Desktop\\my-cv.docx")

        

class CvMaker(MDApp):
    def build(self):
        Window.size = (600, 620)
        Window.clearcolor = "#22438646"
        self.theme_cls.primary_palette = "White"
        self.theme_cls.theme_style = "Dark"
        return Builder.load_file("C:\\Users\\AUTHENTIC PLUS STORE\\Desktop\\goals for the holidays\\sleepy panda\\cv maker\\cv.kv")

CvMaker().run()
